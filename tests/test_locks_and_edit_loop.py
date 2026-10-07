"""
Manual locks and the human-review edit/re-solve loop.

Covers:
  (a) a valid lock actually holds in a solved schedule, in both the
      force-in and force-out directions
  (b) an invalid lock (e.g. an assistant-only person locked into a tech
      role) is caught by validate_locks() before it reaches build_model()
  (c) the full pause -> edit -> re-solve -> pause -> approve cycle through
      the LangGraph pipeline, including falling back to the last good
      schedule when an edit makes the week infeasible, the pipeline keeping
      only edits that solved, removing an edit, and saving the approved
      schedule
"""

import openpyxl
from helpers import full_week_availability
from langgraph.checkpoint.sqlite import SqliteSaver
from langgraph.types import Command

from scheduler.db import crud
from scheduler.pipeline.graph import REVIEW, ROSTER_CONFIRM, build_graph
from scheduler.solver.locks import validate_locks
from scheduler.solver.model_input import LockedAssignment, Person, SolverInput
from scheduler.solver.solve import solve as run_solver

DAY_COLUMNS = {"Mon": "C", "Tue": "E", "Wed": "G", "Thu": "I", "Fri": "K", "Sat": "M", "Sun": "O"}


def make_roster() -> list[Person]:
    roster = [Person(i, f"Tech{i}", "tech_only", 3, 1, 15, full_week_availability()) for i in range(1, 11)]
    roster += [Person(i, f"Asst{i}", "assistant_only", 3, 1, 15, full_week_availability()) for i in range(11, 21)]
    return roster


def write_submission_files(roster: list[Person], tmp_path) -> list[str]:
    """Minimal template-shaped xlsx files - parse_xlsx() only reads fixed cells."""
    paths = []
    for p in roster:
        path = tmp_path / f"{p.name}.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["C2"] = p.name
        ws["J2"] = p.hours_requested
        for day, col in DAY_COLUMNS.items():
            for i, is_avail in enumerate(p.availability[day]):
                if is_avail:
                    ws[f"{col}{6 + i}"] = "X"
        wb.save(path)
        paths.append(str(path))
    return paths


def test_valid_locks_hold():
    roster = make_roster()
    data = SolverInput(people=roster)
    baseline = run_solver(data, time_limit_seconds=20)
    assert baseline["feasible"], f"baseline expected feasible, got {baseline['status']}"

    # Force-out: take a real baseline assignment and lock that person OUT of it.
    victim = baseline["assignments"][0]
    force_out = LockedAssignment(
        person_id=victim["person_id"],
        day=victim["day"], slot=victim["slot"], role=victim["role"], value=False,
    )
    # Force-in: a tech-capable person into a slot they're available for.
    tech2 = next(p for p in roster if p.name == "Tech2")
    force_in = LockedAssignment(person_id=tech2.id, day="Wed", slot=10, role="tech", value=True)

    data.locked_assignments = [force_out, force_in]
    assert validate_locks(data) == []

    result = run_solver(data, time_limit_seconds=20)
    assert result["feasible"], f"expected feasible with locks applied, got {result['status']}"

    assert not any(
        a["name"] == victim["name"] and a["day"] == victim["day"]
        and a["slot"] == victim["slot"] and a["role"] == victim["role"]
        for a in result["assignments"]
    ), "force-out lock did not hold"
    assert any(
        a["name"] == "Tech2" and a["day"] == "Wed" and a["slot"] == 10 and a["role"] == "tech"
        for a in result["assignments"]
    ), "force-in lock did not hold"


def test_invalid_lock_caught_before_build_model():
    roster = make_roster()
    assistant = next(p for p in roster if p.role_weighting == "assistant_only")
    bad_lock = LockedAssignment(person_id=assistant.id, day="Mon", slot=5, role="tech", value=True)

    data = SolverInput(people=roster, locked_assignments=[bad_lock])
    conflicts = validate_locks(data)
    assert conflicts, "expected a capability conflict for an assistant-only person locked into tech"
    assert "tech" in conflicts[0].lower()
    assert "Mon 10:30" in conflicts[0], "conflicts should name the slot by its clock time"

    # Without validation, the solver must not quietly honor the lock either.
    assert not run_solver(data, time_limit_seconds=5)["feasible"]


def test_second_evening_tech_lock_caught():
    tech1, tech2 = make_roster()[:2]
    locks = [LockedAssignment(p.id, "Tue", 20, "tech", True) for p in (tech1, tech2)]
    data = SolverInput(people=[tech1, tech2], locked_assignments=locks)

    assert validate_locks(data) == ["Tue 18:00: 2 people forced into the tech role (Tech1, Tech2), but it only takes 1"]
    # Two in the tech role at 4:00 is fine.
    data.locked_assignments = [LockedAssignment(p.id, "Tue", 16, "tech", True) for p in (tech1, tech2)]
    assert validate_locks(data) == []


def test_full_edit_resolve_cycle(tmp_path):
    roster = make_roster()
    paths = write_submission_files(roster, tmp_path)

    # The graph builds its solver input from the DB roster, so seed complete
    # rows - roster_completeness_check then passes with nothing to fix.
    db_id_by_name = {
        p.name: crud.add_person(
            p.name, f"I{p.id}", p.role_weighting, p.experience_rating, p.proximity, p.hours_requested
        ).id
        for p in roster
    }

    with SqliteSaver.from_conn_string(str(tmp_path / "checkpoints.db")) as checkpointer:
        app = build_graph(checkpointer)
        config = {"configurable": {"thread_id": "edit-loop-test"}}

        def edit(**changes):
            return app.invoke(Command(resume={"decision": "edit", **changes}), config=config)

        def accepted_locks(result):
            return result["__interrupt__"][0].value["locks"]

        result = app.invoke({"submission_file_paths": paths}, config=config)
        payload = result["__interrupt__"][0].value
        assert payload["kind"] == ROSTER_CONFIRM

        decisions = [{"index": c["index"], "action": "confirm"} for c in payload["candidates"]]
        result = app.invoke(Command(resume={"decisions": decisions}), config=config)
        assert result["__interrupt__"][0].value["kind"] == REVIEW

        # A valid edit loops back to solve and the lock shows up in the schedule.
        lock = {"person_id": db_id_by_name["Tech2"], "day": "Thu", "slot": 12, "role": "tech", "value": True}
        result = edit(add_locks=[lock])
        assert not result.get("lock_conflicts")
        assert accepted_locks(result) == [lock]
        assert any(
            a["name"] == "Tech2" and a["day"] == "Thu" and a["slot"] == 12 and a["role"] == "tech"
            for a in result["solve_result"]["assignments"]
        )
        good_solve_result = result["solve_result"]

        # A contradictory lock is rejected before solving, and isn't kept.
        bad = {"person_id": db_id_by_name["Asst11"], "day": "Mon", "slot": 4, "role": "tech", "value": True}
        result = edit(add_locks=[bad])
        assert result.get("lock_conflicts")
        assert result["solve_result"] == good_solve_result
        assert accepted_locks(result) == [lock]

        # Force every tech OUT of one weekend slot. Each lock is individually
        # valid, so validate_locks() passes, but weekend coverage (hard) becomes
        # unsatisfiable. The pipeline must fall back to the last good schedule
        # and pause at human_review again rather than dead-ending at explain.
        breaking_locks = [
            {"person_id": db_id_by_name[f"Tech{i}"], "day": "Sat", "slot": 10, "role": "tech", "value": False}
            for i in range(1, 11)
        ]
        result = edit(add_locks=breaking_locks)
        assert not result.get("lock_conflicts")
        assert result.get("infeasibility_gaps")
        assert result["solve_result"] == good_solve_result
        # The failed edit is dropped: the accepted locks are still the ones
        # the schedule on screen was solved with.
        assert accepted_locks(result) == [lock]

        # The next edit builds on the accepted locks, not the failed ones.
        second = {"person_id": db_id_by_name["Asst11"], "day": "Mon", "slot": 4, "role": "assistant", "value": True}
        result = edit(add_locks=[second])
        assert not result.get("infeasibility_gaps") and not result.get("lock_conflicts")
        assert accepted_locks(result) == [lock, second]

        # An edit can remove a lock again.
        result = edit(remove_locks=[lock])
        assert not result.get("infeasibility_gaps") and not result.get("lock_conflicts")
        assert accepted_locks(result) == [second]

        # Approving saves the schedule, and the locks it was solved with,
        # outside the pipeline's own state.
        final = app.invoke(Command(resume={"decision": "approved"}), config=config)
        assert final["final_schedule"] is not None
        saved = crud.latest_schedule()
        assert saved.id == final["schedule_id"]
        assert saved.result == final["final_schedule"]
        assert saved.locks == [second]
