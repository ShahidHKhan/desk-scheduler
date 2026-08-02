"""
Parses a filled-in Blank_Schedule.xlsx submission into an AvailabilitySubmission.

Cell layout confirmed by inspecting the actual template
(/mnt/project/Blank_Schedule.xlsx):
  - Name:            C2 (merged C2:H2)
  - Hours requested:  J2 (merged J2:K2)
  - Day header row:   row 4 (Mon=C, Tue=E, Wed=G, Thu=I, Fri=K, Sat=M, Sun=O)
  - Time slot rows:   rows 6-30 (25 rows), one per TIME_SLOTS entry
  - Each day's data column is the LEFT cell of a merged 2-column pair
    (e.g. Monday's slots live in column C, with D merged alongside it)
"""

import openpyxl

from schema import DAYS, TIME_SLOTS, AvailabilitySubmission, hours_range_warning, is_available_value

# Column letter that holds each day's availability data (left cell of the merge)
DAY_COLUMNS = {
    "Mon": "C", "Tue": "E", "Wed": "G", "Thu": "I",
    "Fri": "K", "Sat": "M", "Sun": "O",
}

NAME_CELL = "C2"
HOURS_CELL = "J2"
FIRST_SLOT_ROW = 6  # row for TIME_SLOTS[0] == "08:00"


def parse_xlsx(path: str) -> AvailabilitySubmission:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active  # the template only has one sheet ("Semester")

    name = _clean(ws[NAME_CELL].value) or "UNKNOWN"
    hours_raw = ws[HOURS_CELL].value
    hours_requested = _parse_hours(hours_raw)

    warnings = []
    if name == "UNKNOWN" or name.strip("* ").lower() in ("your name", ""):
        warnings.append("Name cell looks unfilled (still has placeholder text)")
    if hours_requested is None:
        warnings.append(f"Could not parse hours_requested from {HOURS_CELL!r} = {hours_raw!r}")
    else:
        range_warning = hours_range_warning(hours_requested)
        if range_warning:
            warnings.append(range_warning)

    availability: dict[str, list[bool]] = {}
    for day in DAYS:
        col = DAY_COLUMNS[day]
        flags = []
        for i in range(len(TIME_SLOTS)):
            row = FIRST_SLOT_ROW + i
            cell_value = ws[f"{col}{row}"].value
            flags.append(is_available_value(cell_value))
        availability[day] = flags

    return AvailabilitySubmission(
        name=name,
        hours_requested=hours_requested,
        availability=availability,
        source_file=path,
        parser_used="xlsx",
        warnings=warnings,
    )


def _clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_hours(value) -> int | None:
    if value is None:
        return None
    try:
        return int(float(value))
    except (ValueError, TypeError):
        return None


if __name__ == "__main__":
    import sys

    result = parse_xlsx(sys.argv[1] if len(sys.argv) > 1 else "/mnt/project/Blank_Schedule.xlsx")
    print(result)
    for day in DAYS:
        print(day, result.available_slots(day))
