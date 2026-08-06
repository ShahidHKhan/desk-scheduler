"""
Synthetic self-test for Phase 5's lock support and edit/re-solve loop.

Run directly:
    python test_phase5.py

Proves:
  (a) a valid lock actually holds in a solved schedule, both force-in and
      force-out directions
  (b) an invalid lock (e.g. an assistant-only person locked into a tech
      role) is caught by validate_locks() and never reaches build_model()
  (c) the full pause -> edit -> re-solve (via the new human_review -> solve
      loop edge) -> pause -> approve cycle works end-to-end via direct
      Command(resume=...) calls
"""

import shutil
import tempfile

import openpyxl
from langgraph.checkpoint.sqlite import SqliteSaver
from langgraph.types import Command

import crud
import database
from build_model import build_model
from graph import build_graph
from locks import validate_locks
from model_input import LockedAssignment, Person, SolverInput
from solve import solve as run_solver


def full_week_availability():
    weekday = [True] * 24 + [False]
    fri = [True] * 18 + [False] * 7
    weekend = [False] * 8 + [True] * 10 + [False] * 7
    return {"Mon": weekday, "Tue": weekday, "Wed": weekday, "Thu": weekday,
            "Fri": fri, "Sat": weekend, "Sun": weekend}


def make_roster():
    roster = [Person(i, f"Tech{i}", "tech_only", 3, 1, 15, full_week_availability()) for i in range(1, 11)]
    roster += [Person(i, f"Asst{i}", "assistant_only", 3, 1, 15, full_week_availability()) for i in range(11, 21)]
    return roster


def test_valid_locks_hold():
    print("=== (a) valid locks hold, both directions ===")
    roster = make_roster()
    data = SolverInput(people=roster)
    baseline = run_solver(data, time_limit_seconds=20)
    assert baseline["feasible"], f"baseline expected feasible, got {baseline['status']}"

    # Force-out: pick a real baseline assignment and lock that person OUT of it.
    victim = baseline["assignments"][0]
    force_out = LockedAssignment(
        person_id=next(p.id for p in roster if p.name == victim["person"]),
        day=victim["day"], slot=victim["slot"], role=victim["role"], value=False,
    )

    # Force-in: pick a tech-capable person, a slot they're available for but
    # weren't necessarily assigned, and force them into it.
    tech1 = next(p for p in roster if p.name == "Tech2")
    force_in = LockedAssignment(person_id=tech1.id, day="Wed", slot=10, role="tech", value=True)

    data.locked_assignments = [force_out, force_in]
    assert validate_locks(data) == [], "expected these locks to be valid"

    result = run_solver(data, time_limit_seconds=20)
    assert result["feasible"], f"expected feasible with locks applied, got {result['status']}"

    still_present = any(
        a["person"] == victim["person"] and a["day"] == victim["day"]
        and a["slot"] == victim["slot"] and a["role"] == victim["role"]
        for a in result["assignments"]
    )
    assert not still_present, "force-out lock did not hold"

    forced_in_present = any(
        a["person"] == "Tech2" and a["day"] == "Wed" and a["slot"] == 10 and a["role"] == "tech"
        for a in result["assignments"]
    )
    assert forced_in_present, "force-in lock did not hold"
    print("OK: force-out and force-in locks both held.\n")


def test_invalid_lock_caught_before_build_model():
    print("=== (b) invalid lock caught by validate_locks(), never reaches build_model() ===")
    roster = make_roster()
    assistant = next(p for p in roster if p.role_weighting == "assistant_only")
    bad_lock = LockedAssignment(person_id=assistant.id, day="Mon", slot=5, role="tech", value=True)

    data = SolverInput(people=roster, locked_assignments=[bad_lock])
    conflicts = validate_locks(data)
    assert conflicts, "expected a capability conflict for an assistant-only person locked into tech"
    assert "tech" in conflicts[0].lower()
    print(f"OK: caught -> {conflicts[0]}")

    # Confirm build_model() itself would blow up on this (KeyError or an
    # unsatisfiable model) if validation were skipped - i.e. this really is
    # locks.py's job, not something build_model quietly tolerates.
    try:
        model, variables = build_model(data)
        solver_result = run_solver(data, time_limit_seconds=5)
        assert not solver_result["feasible"], (
            "expected an unenforceable tech lock on an assistant-only person to make the model infeasible"
        )
        print("OK: confirmed build_model()/solve() do not silently honor this lock either.\n")
    except KeyError:
        print("OK: build_model() raised KeyError on the bad lock (also acceptable).\n")


def _make_synthetic_xlsx_paths(roster, tmpdir):
    paths = []
    day_cols = {"Mon": "C", "Tue": "E", "Wed": "G", "Thu": "I", "Fri": "K", "Sat": "M", "Sun": "O"}
    for p in roster:
        path = f"{tmpdir}/{p.name}.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["C2"] = p.name
        ws["J2"] = p.hours_requested
        avail = full_week_availability()
        for day, col in day_cols.items():
            for i, is_avail in enumerate(avail[day]):
                if is_avail:
                    ws[f"{col}{6 + i}"] = "X"
        wb.save(path)
        paths.append(path)
    return paths


def test_full_edit_resolve_cycle():
    print("=== (c) pause -> edit -> re-solve -> pause -> approve, end-to-end ===")
    roster = make_roster()
    tmpdir = tempfile.mkdtemp()
    paths = _make_synthetic_xlsx_paths(roster, tmpdir)

    # Seed the DB-backed roster (the graph now builds its solver input from
    # crud.list_roster(), not from a "roster" key passed into invoke()) -
    # complete rows so roster_completeness_check passes with nothing to fix.
    database.init_db()
    db_id_by_name = {}
    for p in roster:
        db_person = crud.add_person(
            p.name, f"I{p.id}", p.role_weighting, p.experience_rating, p.proximity, p.hours_requested
        )
        db_id_by_name[p.name] = db_person.id

    try:
        with SqliteSaver.from_conn_string(f"{tmpdir}/checkpoints.db") as checkpointer:
            app = build_graph(checkpointer)
            config = {"configurable": {"thread_id": "phase5-test"}}

            result = app.invoke({"submission_file_paths": paths}, config=config)
            assert "__interrupt__" in result, f"expected a pause at roster_confirm, got {result.get('solve_status')}"
            interrupt_payload = result["__interrupt__"][0].value
            assert "candidates" in interrupt_payload, "expected the first pause to be roster_confirm"
            print("OK: paused at roster_confirm.")

            # Every submission's name exactly matches a roster row we just
            # seeded - confirm every suggested match as-is.
            decisions = [{"index": c["index"], "action": "confirm"} for c in interrupt_payload["candidates"]]
            result = app.invoke(Command(resume={"decisions": decisions}), config=config)
            assert "__interrupt__" in result, f"expected a pause at human_review, got {result.get('solve_status')}"
            assert "candidates" not in result["__interrupt__"][0].value, (
                "roster was already complete - should not re-pause at roster_confirm or the completeness gate"
            )
            print("OK: roster confirmed, roster complete, paused at human_review (first solve).")

            _run_edit_resolve_assertions(app, config, db_id_by_name)
    finally:
        for person_id in db_id_by_name.values():
            crud.delete_person(person_id)

    shutil.rmtree(tmpdir, ignore_errors=True)


def _run_edit_resolve_assertions(app, config, db_id_by_name):
    # Force a specific tech onto a specific slot via an edit, and loop back to solve.
    tech2_id = db_id_by_name["Tech2"]
    lock = {"person_id": tech2_id, "day": "Thu", "slot": 12, "role": "tech", "value": True}
    result = app.invoke(Command(resume={"decision": "edit", "locked_assignments": [lock]}), config=config)
    assert not result.get("lock_conflicts"), f"unexpected lock conflict: {result.get('lock_conflicts')}"
    assert "__interrupt__" in result, "expected the edit to loop back to solve and pause again at human_review"
    print("OK: edit accepted, looped solve -> human_review, paused again.")

    forced = any(
        a["person"] == "Tech2" and a["day"] == "Thu" and a["slot"] == 12 and a["role"] == "tech"
        for a in result["solve_result"]["assignments"]
    )
    assert forced, "the edit's lock should be reflected in the re-solved schedule"
    good_solve_result = result["solve_result"]
    print("OK: re-solved schedule reflects the edit.")

    # Now the case that a naive "loop back only on lock_conflicts" fix
    # would miss: force every tech-capable person OUT of one WEEKEND slot.
    # Each individual lock is valid (a force-out can't conflict with
    # availability/capability), so validate_locks() passes clean - but the
    # weekend's exactly-1-tech coverage requirement (still a hard
    # constraint - see build_model.py) becomes unsatisfiable, so the
    # CP-SAT solve itself goes INFEASIBLE. This must fall back to the last
    # good schedule and re-pause at human_review, not dead-end at explain.
    # (Note: this used to force a WEEKDAY slot instead, but weekday
    # coverage is now a soft/penalized constraint - forcing 10 techs out
    # of one weekday slot no longer breaks feasibility, it just gets
    # reported as a shortfall. Weekend coverage is what's still hard.)
    breaking_locks = [lock] + [
        {"person_id": db_id_by_name[f"Tech{i}"], "day": "Sat", "slot": 10, "role": "tech", "value": False}
        for i in range(1, 11)
    ]
    result = app.invoke(
        Command(resume={"decision": "edit", "locked_assignments": breaking_locks}), config=config
    )
    assert not result.get("lock_conflicts"), (
        f"each force-out lock is individually valid, expected no lock_conflicts, got {result.get('lock_conflicts')}"
    )
    assert result.get("infeasibility_gaps"), "expected the coverage-breaking edit to surface infeasibility_gaps"
    assert "__interrupt__" in result, (
        "a re-solve failure with a prior good schedule must fall back to human_review, not dead-end at explain"
    )
    assert result["solve_result"] == good_solve_result, (
        "solve_result must still be the last known-good schedule, not None and not the failed attempt"
    )
    print("OK: coverage-breaking (but lock-valid) edit fell back to human_review with the prior schedule intact.")

    # Recover: resend just the good lock, drop the coverage-breaking batch.
    result = app.invoke(Command(resume={"decision": "edit", "locked_assignments": [lock]}), config=config)
    assert not result.get("infeasibility_gaps") and not result.get("lock_conflicts")
    assert "__interrupt__" in result
    print("OK: recovered with a clean edit after the failed one.")

    final = app.invoke(Command(resume={"decision": "approved"}), config=config)
    assert final["final_schedule"] is not None, "expected a final_schedule after approval"
    print("OK: approved, final_schedule produced.\n")


if __name__ == "__main__":
    test_valid_locks_hold()
    test_invalid_lock_caught_before_build_model()
    test_full_edit_resolve_cycle()
    print("All Phase 5 synthetic tests passed.")
